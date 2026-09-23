import io

import pytest

from app.readers import UnsupportedFile, to_text


def test_csv_semicolon_cp1251():
    raw = "код;объём\n1.1;10\n".encode("cp1251")
    out = to_text("a.csv", raw)
    assert "[row 2] 1.1 | 10" in out


def test_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Смета"
    ws.append(["поз", "кол-во"]); ws.append([1, 5])
    buf = io.BytesIO(); wb.save(buf)
    assert "[sheet Смета row 2] 1 | 5" in to_text("s.xlsx", buf.getvalue())


def test_docx():
    import docx
    d = docx.Document(); d.add_paragraph("Привет"); t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "a"; t.cell(0, 1).text = "b"
    buf = io.BytesIO(); d.save(buf)
    out = to_text("d.docx", buf.getvalue())
    assert "[para 1] Привет" in out and "[table 1 row 1] a | b" in out


def test_unsupported():
    with pytest.raises(UnsupportedFile):
        to_text("x.exe", b"MZ")


def test_upload_endpoint():
    import os
    os.environ["MOCK"] = "1"
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    r = c.post("/api/upload", files={"file": ("t.csv", "a;b\n1;2\n".encode(), "text/csv")})
    assert r.status_code == 200, r.text
    assert c.post("/api/upload", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 415

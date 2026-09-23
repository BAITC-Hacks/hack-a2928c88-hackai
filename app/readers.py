"""Превращение загруженных файлов в текст с локаторами, понятными модели и правилу R-QUOTE.

Каждый формат даёт текст с метками:
  PDF  -> "[page N]" перед страницей
  XLSX -> "[sheet S row R]" перед строкой, ячейки через " | "
  CSV  -> "[row R]" перед строкой
  DOCX -> "[para N]" / "[table T row R]"
Модель должна указывать эти метки в evidence.locator.
Сканы PDF без текстового слоя возвращают пустые страницы — тогда нужен OCR (см. TODO ниже).
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

MAX_CHARS = 150_000


class UnsupportedFile(ValueError):
    pass


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFile("Не удалось определить кодировку текста")


def read_pdf(raw: bytes) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            # TODO(ocr): если text пуст — страница-скан; подключить OCR или vision-модель
            parts.append(f"[page {n}]\n{text}")
    return "\n\n".join(parts)


def read_xlsx(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                parts.append(f"[sheet {ws.title} row {r}] " + " | ".join(cells).rstrip(" |"))
    return "\n".join(parts)


def read_csv(raw: bytes) -> str:
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return "\n".join(f"[row {r}] " + " | ".join(row) for r, row in enumerate(rows, 1) if any(row))


def read_docx(raw: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(raw))
    parts = [f"[para {n}] {p.text}" for n, p in enumerate(d.paragraphs, 1) if p.text.strip()]
    for t, table in enumerate(d.tables, 1):
        for r, row in enumerate(table.rows, 1):
            parts.append(f"[table {t} row {r}] " + " | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


READERS = {
    ".pdf": read_pdf,
    ".xlsx": read_xlsx, ".xlsm": read_xlsx,
    ".csv": read_csv, ".tsv": read_csv,
    ".docx": read_docx,
}
TEXT_EXT = {".txt", ".md", ".json", ".eml", ".log", ".xml", ".html", ".htm", ""}


def to_text(filename: str, raw: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in READERS:
        text = READERS[ext](raw)
    elif ext in TEXT_EXT:
        text = _decode(raw)
    else:
        raise UnsupportedFile(f"Формат {ext} пока не поддерживается: pdf, xlsx, csv, docx, txt, md, json, eml")
    if not text.strip():
        raise UnsupportedFile("Файл не содержит извлекаемого текста (скан? нужен OCR)")
    return text[:MAX_CHARS]
